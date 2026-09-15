"""排班结果的结构化表格构建与导出（Excel / CSV）"""

from __future__ import annotations

import io
from collections import defaultdict

import pandas as pd

from .database import Assignment
from .parser import BLOCK_LABELS, WEEKDAY_LABELS


def build_detail_df(assignments: list[Assignment], per_slot: int = 1) -> pd.DataFrame:
    """明细表：每行一个值班任务（周次 x 星期 x 时段），多人在岗合并显示"""
    grouped: dict[tuple[int, int, int], list[str]] = defaultdict(list)
    for a in assignments:
        grouped[(a.week, a.weekday, a.block)].append(a.member_name)

    rows = []
    for (week, weekday, block), names in sorted(grouped.items()):
        rows.append({
            "周次": f"第{week}周",
            "星期": WEEKDAY_LABELS[weekday],
            "值班时段": BLOCK_LABELS[block],
            "值班人": "、".join(sorted(names)),
            "人数": len(names),
        })
    return pd.DataFrame(rows, columns=["周次", "星期", "值班时段", "值班人", "人数"])


def build_pivot_df(assignments: list[Assignment]) -> pd.DataFrame:
    """透视表：行=周次x星期，列=值班时段（多人用顿号连接），适合界面展示与打印"""
    grouped: dict[tuple[int, int, int], list[str]] = defaultdict(list)
    for a in assignments:
        grouped[(a.week, a.weekday, a.block)].append(a.member_name)

    blocks = sorted({b for (_, _, b) in grouped})
    weeks = sorted({w for (w, _, _) in grouped})
    weekdays = sorted({d for (_, d, _) in grouped})

    index = [(w, d) for w in weeks for d in weekdays]
    data = {}
    for b in blocks:
        col = []
        for (w, d) in index:
            names = grouped.get((w, d, b))
            col.append("、".join(sorted(names)) if names else "—")
        data[BLOCK_LABELS[b].split(" ")[0]] = col  # 列名如 "1-2节"

    df = pd.DataFrame(data, index=pd.MultiIndex.from_tuples(
        index, names=["周次", "星期"]))
    df.index = df.index.map(lambda t: (f"第{t[0]}周", WEEKDAY_LABELS[t[1]]))
    return df


def build_stats_df(member_stats: dict[int, dict]) -> pd.DataFrame:
    """值班统计表：每人总次数与值班周分布"""
    rows = [{
        "成员": s["name"],
        "总值班次数": s["total"],
        "值班周次": "、".join(f"第{w}周" for w in s["weeks"]) or "—",
    } for s in member_stats.values()]
    return pd.DataFrame(rows, columns=["成员", "总值班次数", "值班周次"]).sort_values(
        ["总值班次数", "成员"], ascending=[False, True]).reset_index(drop=True)


def export_excel(
    assignments: list[Assignment],
    member_stats: dict[int, dict],
    gaps: list[tuple[int, int, int]],
) -> bytes:
    """导出 Excel：排班总表(透视) + 值班明细 + 值班统计，返回文件字节"""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        pivot = build_pivot_df(assignments)
        pivot.to_excel(writer, sheet_name="排班总表", merge_cells=False)

        detail = build_detail_df(assignments)
        detail.to_excel(writer, sheet_name="值班明细", index=False)

        stats = build_stats_df(member_stats)
        stats.to_excel(writer, sheet_name="值班统计", index=False)

        if gaps:
            gap_df = pd.DataFrame([{
                "周次": f"第{w}周", "星期": WEEKDAY_LABELS[d], "时段": BLOCK_LABELS[b],
            } for w, d, b in gaps])
            gap_df.to_excel(writer, sheet_name="无人可用时段", index=False)

        _beautify(writer)
    return buf.getvalue()


def _beautify(writer: pd.ExcelWriter) -> None:
    """简单美化：列宽自适应、表头加粗"""
    from openpyxl.styles import Font, PatternFill

    for ws in writer.book.worksheets:
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            width = max(
                (len(str(c.value)) + sum(1 for ch in str(c.value) if '\u4e00' <= ch <= '\u9fff')
                 for c in column_cells if c.value is not None),
                default=8,
            )
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(
                width + 4, 40)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")


def export_csv(assignments: list[Assignment]) -> bytes:
    """导出 CSV（UTF-8 BOM，Excel 可直接打开）"""
    return build_detail_df(assignments).to_csv(index=False).encode("utf-8-sig")
