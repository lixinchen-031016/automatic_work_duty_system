"""Excel 导出布局：按内容自动设置列宽、换行与行高。"""

from __future__ import annotations

import unicodedata
from copy import copy
from math import ceil

from openpyxl.utils import get_column_letter

MAX_COLUMN_WIDTH = 42.0
MIN_COLUMN_WIDTH = 10.0
LINE_HEIGHT = 16.0
ROW_PADDING = 5.0
MAX_ROW_HEIGHT = 409.0


def excel_text_width(text: str) -> int:
    """按 Excel 字符宽近似计算文本宽度，中日韩字符按 2 个字符计。"""
    return sum(
        2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
        for char in text
    )


def excel_line_count(text: str, column_width: float) -> int:
    """估算文本在指定列宽下换行后的行数。"""
    usable_width = max(4.0, column_width - 2.0)
    return sum(
        max(1, ceil(excel_text_width(line) / usable_width))
        for line in text.split("\n")
    )


def fit_excel_layout(
    ws,
    *,
    min_widths: dict[int, float] | None = None,
    max_column_width: float = MAX_COLUMN_WIDTH,
    min_row_height: float = 22.0,
    header_rows: int = 0,
) -> None:
    """按单元格内容设置列宽和行高，并开启自动换行。

    调用方只需在写入完单元格后调用一次。日期、课程、人名等长文本会自动换行，
    行高按最宽内容估算，打开 Excel 后无需手工调整单元格大小。
    """
    min_widths = min_widths or {}
    last_column = max(ws.max_column, max(min_widths, default=0))
    widths: dict[int, float] = {
        column_index: float(min_widths.get(column_index, MIN_COLUMN_WIDTH))
        for column_index in range(1, last_column + 1)
    }

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or str(cell.value) == "":
                continue
            widths[cell.column] = max(
                widths[cell.column],
                min(
                    max(
                        excel_text_width(line)
                        for line in str(cell.value).split("\n")
                    ) + 2.0,
                    max_column_width,
                ),
            )

    for column_index, width in widths.items():
        ws.column_dimensions[get_column_letter(column_index)].width = width

    for row in ws.iter_rows():
        required_lines = 1
        for cell in row:
            if cell.value is None or str(cell.value) == "":
                continue
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment
            required_lines = max(
                required_lines,
                excel_line_count(str(cell.value), widths[cell.column]),
            )

        row_height = max(
            min_row_height,
            required_lines * LINE_HEIGHT + ROW_PADDING,
        )
        if header_rows and row[0].row <= header_rows:
            row_height = max(row_height, 24.0)
        ws.row_dimensions[row[0].row].height = min(MAX_ROW_HEIGHT, row_height)
